"""Tests for supply_chain_processor.py.

Run from the repository root:  python -m unittest discover tests
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import supply_chain_processor as scp  # noqa: E402

BASIL = {'code': '321601', 'baseQty': 24, 'cbmPerUnit': 0.00128, 'capacity20ft': 1000, 'capacity40ft': 2100}


class StatusTests(unittest.TestCase):
    def test_thresholds_match_dashboard(self):
        self.assertEqual(scp.get_status(None), 'NO SALES')
        self.assertEqual(scp.get_status(9.9), 'CRITICAL')
        self.assertEqual(scp.get_status(10), 'LOW')
        self.assertEqual(scp.get_status(19.9), 'LOW')
        self.assertEqual(scp.get_status(20), 'OPTIMAL')
        self.assertEqual(scp.get_status(90), 'OPTIMAL')
        self.assertEqual(scp.get_status(90.1), 'HIGH')


class ContainerTests(unittest.TestCase):
    def test_cbm_per_carton_uses_pieces_per_carton(self):
        self.assertAlmostEqual(scp.cbm_per_carton(BASIL), 0.03072)
        self.assertAlmostEqual(scp.cbm_per_carton({'cbmPerUnit': 0.1}), 0.1)

    def test_fills_40ft_then_remainder_in_20ft(self):
        plan = scp.plan_containers(2100 + 900, BASIL)
        self.assertEqual((plan['containers40ft'], plan['containers20ft']), (1, 1))
        self.assertEqual(plan['label'], '1 x 40ft + 1 x 20ft')
        self.assertFalse(plan['estimated'])

    def test_remainder_too_big_for_20ft_takes_another_40ft(self):
        plan = scp.plan_containers(7850, BASIL)
        self.assertEqual((plan['containers40ft'], plan['containers20ft']), (4, 0))
        self.assertAlmostEqual(plan['fill'], 0.935, places=3)

    def test_preferred_size_is_respected(self):
        plan = scp.plan_containers(2500, BASIL, preferred='20ft')
        self.assertEqual((plan['containers40ft'], plan['containers20ft']), (0, 3))

    def test_capacity_estimated_from_cbm_when_missing(self):
        item = {'baseQty': 4, 'cbmPerUnit': 0.025}  # 0.1 m3 per carton -> 280 per 20ft, 580 per 40ft
        plan = scp.plan_containers(600, item)
        self.assertEqual((plan['containers40ft'], plan['containers20ft']), (1, 1))
        self.assertTrue(plan['estimated'])

    def test_no_capacity_or_cbm_gives_no_plan(self):
        self.assertIsNone(scp.plan_containers(100, {'baseQty': 24}))


def build_workbook(path):
    """A small workbook in the MISB layout (title rows above the headers)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'ItemData'
    ws.append(['Item Code', 'Item Name', 'BaseQty', '40ft', '20ft', 'cbm', 'Supplier Name', 'VendorCode', 'MOQ', 'Transit Time(Days)'])
    ws.append([321601, 'Basil Mango', 24, 2100, 1000, 0.00128, 'YU DAT BS', 'V20183', 450, 80])
    ws.append([120405, 'Atta', 4, None, None, 0.025, 'Local A', 'V20185', 1000, 7])
    ws.append([999, 'Never sold', 24, 2100, 1000, 0.00128, 'YU DAT BS', 'V20183', 450, 80])

    ws = wb.create_sheet('MisbWH')
    ws.append(['Warehouse report'])
    ws.append(['Item Code', 'Item Name', 'Qty(Ctn)', 'Qty(Pcs)'])
    ws.append([321601.0, 'Basil Mango', 2500, 60000])  # float code must still match 321601
    ws.append([120405, 'Atta', 2472, 9888])
    ws.append([999, 'Never sold', 10, 240])

    ws = wb.create_sheet('LocalItemReport')
    for _ in range(3):
        ws.append(['title'])
    header = ['Code']
    for _ in range(12):
        header += ['BAL', 'INC', 'DEC']
    ws.append(header)

    def sales_row(code, recent, older):
        row = [code]
        for month in range(12):
            row += [0, 0, recent if month >= 9 else older]
        return row

    # Last 3 months: 2,550/month (85/day) and 5,400/month (180/day). Older months must be ignored.
    ws.append(sales_row(321601, 2550, 99999))
    ws.append(sales_row(120405, 5400, 99999))
    wb.save(path)


class EndToEndTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        build_workbook(cls.dir / 'test.xlsx')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def run_processor(self):
        p = scp.SupplyChainProcessor()
        p.read_excel_file(self.dir / 'test.xlsx')
        p.process_item_master()
        p.process_current_inventory()
        p.calculate_average_daily_sales(months=3)
        p.calculate_po_quantities(inventory_days=30, safety_stock_days=7)
        return p

    def test_average_uses_only_last_three_months(self):
        stats = self.run_processor().sales_stats
        self.assertEqual(stats['321601']['avgDaily'], 85)
        self.assertEqual(stats['120405']['avgDaily'], 180)
        self.assertNotIn('999', stats)

    def test_po_matches_dashboard_example(self):
        pos = {po['itemCode']: po for po in self.run_processor().po_calculations}
        basil = pos['321601']
        self.assertEqual(basil['poQty'], 7850)
        self.assertEqual(basil['containerPlan'], '4 x 40ft')
        self.assertAlmostEqual(basil['cbmTotal'], 241.15)
        self.assertEqual(basil['priority'], 'OPTIMAL')

        atta = pos['120405']
        self.assertEqual(atta['poQty'], 5528)
        self.assertTrue(atta['capacityEstimated'])
        self.assertEqual(atta['priority'], 'LOW')
        self.assertNotIn('999', pos)  # no sales -> no PO

    def test_cli_writes_json_and_csv(self):
        out = self.dir / 'out'
        csv_path = self.dir / 'po.csv'
        code = scp.main([str(self.dir / 'test.xlsx'), '--output-dir', str(out), '--csv', str(csv_path)])
        self.assertEqual(code, 0)
        stats = json.loads((out / 'sales_stats.json').read_text())
        self.assertEqual(stats['321601']['totalSales'], 7650)
        self.assertTrue(csv_path.exists())

    def test_cli_reports_failure(self):
        self.assertEqual(scp.main([str(self.dir / 'missing.xlsx'), '--output-dir', str(self.dir / 'x')]), 1)


if __name__ == '__main__':
    unittest.main()
